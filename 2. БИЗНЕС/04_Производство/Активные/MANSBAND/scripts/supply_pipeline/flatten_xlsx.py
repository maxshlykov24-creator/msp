"""Сырой xlsx поставки → flat-строки (1 размер = 1 строка)."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

from config import KIT_MAP
from expand import _clean, _qty, norm_kit, resolve_kit
from sostav import normalize_sostav


def load_spr_kit_parts(path: Path) -> Dict[str, List[str]]:
    """Лист «Справочник»: блок = вид комплекта, ниже части (пиджак, брюки, жилет)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet_name = next((n for n in wb.sheetnames if "правочник" in n.lower()), None)
    if not sheet_name:
        return {}
    ws = wb[sheet_name]
    lines: List[str] = []
    for row in ws.iter_rows(max_col=1, values_only=True):
        lines.append(" ".join(_clean(row[0]).replace("\xa0", " ").split()))
    out: Dict[str, List[str]] = {}
    i = 0
    while i < len(lines):
        if not lines[i]:
            i += 1
            continue
        header = lines[i]
        i += 1
        parts: List[str] = []
        while i < len(lines) and lines[i]:
            parts.append(lines[i])
            i += 1
        if parts:
            out[header] = parts
    return out


def kit_spr_mismatches(path: Path, kit_raws: List[str]) -> List[str]:
    """
    Если в файле поставки вид есть в «Справочнике», части должны совпасть с KIT_MAP.
    Расхождение — стоп, в МС не пишем. Так 28.08 двубортный оверсайз уехал
    в «пиджак оверсайз с накладными карманами».
    """
    spr = load_spr_kit_parts(path)
    if not spr:
        return []
    by_canon: Dict[str, Tuple[str, List[str]]] = {}
    for header, parts in spr.items():
        by_canon.setdefault(resolve_kit(header), (header, parts))
    errors: List[str] = []
    seen = set()
    for raw in kit_raws:
        canon = resolve_kit(raw)
        hit = by_canon.get(canon)
        if hit is None:
            continue
        header, parts = hit
        expected = KIT_MAP.get(canon)
        if sorted(expected or []) == sorted(parts):
            continue
        key = (norm_kit(header), tuple(sorted(parts)), tuple(sorted(expected or [])))
        if key in seen:
            continue
        seen.add(key)
        errors.append(
            f"вид {raw!r} по справочнику файла → {parts}, "
            f"в KIT_MAP ({canon}) → {expected}. В МС не пишем."
        )
    return errors


FLAT_HEADER = [
    "num",
    "kit",
    "brand",
    "model",
    "rost",
    "art",
    "sostav",
    "color",
    "uzor",
    "kroy",
    "size",
    "qty",
    "note",  # Коробка / пометки (разный узор по частям комплекта и т.п.)
]


def parse_supply_xlsx(path: Path) -> Tuple[List[List[Any]], Dict[str, Any]]:
    """
    Парсит лист вида «Поставка …»: блоки модель + сетка размеров.
    Факт qty = сумма размеров (ячейка «Всего» только для контроля).
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    # первый лист, не «Справочник»
    sheet_name = next(n for n in wb.sheetnames if "правочник" not in n.lower())
    ws = wb[sheet_name]

    flat: List[List[Any]] = [FLAT_HEADER]
    models: List[Dict[str, Any]] = []
    i = 1
    while i <= ws.max_row:
        row = [_clean(ws.cell(i, c).value) for c in range(1, 12)]
        # заголовок блока
        if row[0].lower() == "номер" or (row[2].lower() == "вид" and row[3].lower().startswith("произв")):
            i += 1
            continue
        # строка модели: номер числовой
        num = row[0]
        if num.isdigit():
            meta = {
                "num": num,
                "note": row[1],  # Коробка
                "kit_raw": row[2],
                "kit": resolve_kit(row[2]),
                "brand": row[3],
                "model": row[4],
                "rost": row[5],
                "art": row[6],
                "sostav_raw": row[7],
                "sostav": normalize_sostav(row[7]),
                "color": row[8],
                "uzor": row[9],
                "kroy": row[10],
                "sizes": [],
                "total_cell": None,
            }
            i += 1
            # пропускаем шапку Размер/Штук; пар может быть 2 или 3 (размер 70).
            size_pairs = [(1, 2), (3, 4)]
            if i <= ws.max_row:
                hdr = [_clean(ws.cell(i, c).value) for c in range(1, 8)]
                if any(h.lower().startswith("размер") for h in hdr):
                    size_pairs = [
                        (c, c + 1)
                        for c in range(1, 8, 2)
                        if hdr[c - 1].lower().startswith("размер")
                    ] or [(1, 2), (3, 4)]
                    i += 1
            # сетка размеров: пары колонок Размер/Штук.
            # Важно: строка вида [56, 4, 'Всего:', 25] — сначала учесть 56, потом Всего.
            while i <= ws.max_row:
                cells = [
                    (_clean(ws.cell(i, sc).value), ws.cell(i, qc).value)
                    for sc, qc in size_pairs
                ]
                if not any(size for size, _ in cells):
                    break
                first_size = cells[0][0]
                if first_size.lower() == "номер" or (
                    _clean(ws.cell(i, 3).value).lower() == "вид"
                    and not first_size.isdigit()
                    and not first_size.lower().startswith("всего")
                ):
                    break

                def _add_size(size: str, qv: Any) -> None:
                    if not size or size.lower().startswith("всего"):
                        return
                    q = _qty(qv)
                    if q is None:
                        return
                    meta["sizes"].append((size, q))
                    flat.append(
                        [
                            meta["num"],
                            meta["kit"],
                            meta["brand"],
                            meta["model"],
                            meta["rost"],
                            meta["art"],
                            meta["sostav"],
                            meta["color"],
                            meta["uzor"],
                            meta["kroy"],
                            size,
                            q,
                            meta["note"],
                        ]
                    )

                saw_total = False
                for size, qv in cells:
                    if size.lower().startswith("всего"):
                        meta["total_cell"] = _qty(qv)
                        saw_total = True
                        continue
                    _add_size(size, qv)
                i += 1
                if saw_total:
                    break
            fact = sum(q for _, q in meta["sizes"])
            meta["fact_qty"] = fact
            models.append(meta)
            continue
        i += 1

    sum_fact = sum(m["fact_qty"] for m in models)
    sum_total_cells = sum(m["total_cell"] or 0 for m in models)
    mismatches = [
        {
            "num": m["num"],
            "fact": m["fact_qty"],
            "total_cell": m["total_cell"],
        }
        for m in models
        if m["total_cell"] is not None and m["total_cell"] != m["fact_qty"]
    ]
    stats = {
        "sheet": sheet_name,
        "models": len(models),
        "flat_rows": len(flat) - 1,
        "sum_fact_qty": sum_fact,
        "sum_total_cells": sum_total_cells,
        "qty_mismatches": mismatches,
        "kit_spr_mismatches": kit_spr_mismatches(
            path, [m["kit_raw"] for m in models]
        ),
        "model_summaries": [
            {
                "num": m["num"],
                "kit": m["kit"],
                "kit_raw": m["kit_raw"],
                "brand": m["brand"],
                "model": m["model"],
                "art": m["art"],
                "rost": m["rost"],
                "color": m["color"],
                "uzor": m["uzor"],
                "color_uzor": f"{m['color']} / {m['uzor']}",
                "note": m["note"],
                "sostav": m["sostav"],
                "sizes": len(m["sizes"]),
                "fact_qty": m["fact_qty"],
                "total_cell": m["total_cell"],
            }
            for m in models
        ],
    }
    if stats["kit_spr_mismatches"]:
        raise RuntimeError(
            "Справочник файла не совпал с KIT_MAP:\n"
            + "\n".join(stats["kit_spr_mismatches"])
        )
    return flat, stats


parse_xlsx = parse_supply_xlsx


if __name__ == "__main__":
    import json
    import sys

    p = Path(sys.argv[1])
    flat, stats = parse_supply_xlsx(p)
    print(json.dumps(stats, ensure_ascii=False, indent=2))
    print(f"flat_rows={len(flat)-1}")
