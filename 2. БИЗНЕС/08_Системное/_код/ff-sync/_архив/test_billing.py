"""Прогон новой модели денег и выгрузки КиЗ на моках МойСклад.

Проверяем три вещи из созвона: ставка хранения одна на всех, «тариф» позиции —
это сборка в ₽ за штуку, счётчик стартует после фактической приёмки. Плюс что
выгрузка кодов маркировки отдаёт артикул, GTIN и наименование.

Запуск: PYTHONPATH=.. python3 _архив/test_billing.py
"""

import io
import json
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ["FF_DB"] = os.path.join(tempfile.mkdtemp(), "bill.db")
os.environ["MS_TOKEN"] = "test-token"

import db
import net

CALLS = []
ORDER_ID = "PO-1"
PRODUCT_ID = "PROD-1"
SUPPLY_ID = "SUP-1"
supply_applicable = True
supply_deleted = None


class Fake:
    def __init__(self, status=200, data=None):
        self.status_code = status
        self._data = data

    @property
    def content(self):
        return json.dumps(self._data or {}).encode()

    @property
    def text(self):
        return self.content.decode()

    def json(self):
        return self._data

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError("мок отдал %s" % self.status_code)


def fake_req(method, url, headers=None, **kw):
    CALLS.append((method, url, kw.get("json"), kw.get("params")))
    if "/entity/counterparty/" in url:
        if url.endswith("/accounts"):
            return Fake(200, {"meta": {"size": 1}, "rows": [{"accountNumber": "00000000000000000001", "bankName": "Тестовый банк", "bic": "000000001", "correspondentAccount": "00000000000000000002"}]})
        return Fake(200, {"legalTitle": "Тест ООО", "inn": "0000000001", "legalAddress": "Тестовый адрес", "ogrn": "0000000000001"})
    if url.endswith("/entity/purchaseorder/%s" % ORDER_ID):
        return Fake(200, {"id": ORDER_ID, "supplies": [{"meta": {"href": net.MS_BASE + "/entity/supply/" + SUPPLY_ID}}]})
    if url.endswith("/entity/supply/%s/positions" % SUPPLY_ID):
        return Fake(200, {
            "meta": {"size": 1},
            "rows": [{"quantity": 10, "assortment": {"meta": {"href": net.MS_BASE + "/entity/product/" + PRODUCT_ID}}}],
        })
    if url.endswith("/entity/supply/%s" % SUPPLY_ID):
        return Fake(200, {"id": SUPPLY_ID, "applicable": supply_applicable, "deleted": supply_deleted, "moment": "2026-09-01 12:30:00"})
    if url.endswith("/entity/invoiceout"):
        return Fake(200, {"id": "INV-1", "name": "СЧ-1"})
    if "/entity/service" in url:
        return Fake(200, {"rows": [{"id": "SRV-%s" % i, "name": n} for i, n in enumerate(("Хранение на складе", "Приёмка товара", "Сборка заказов"))]})
    return Fake(404, {"message": "мок не знает %s" % url})


net.req = fake_req
db.init_db()

import account
import billing
from export_xlsx import build_marks

client_id = db.insert_client("test", "Тест ООО", "AGENT-1", "ORG-1", "STORE-1")
db.update_client(client_id, tariff_pick=3.0)

# партия завелась сегодня, приёмки ещё нет: счётчик стоять
lot_id = db.insert_lot(
    client_id, PRODUCT_ID, "ART-1", "2000000000019", "02000000000019", "Ремень кожаный",
    "Не маркируется", 2.0, 10, "2026-09-01T09:00:00", ORDER_ID, pick_rate=None, dims="10x10x20",
)

rows = account.lot_rows(only_open=False)
row = rows[0]
assert row["accepted"] == "", "до приёмки даты быть не должно: %r" % row["accepted"]
assert row["bill_days"] == 0 and row["storage"] == 0, "непринятая партия не должна ничего стоить: %r" % row
assert row["rate"] == 0.15, "ставка хранения общая: %r" % row["rate"]
assert row["pick_rate"] == 3.0, "сборка берётся у клиента, когда у позиции пусто: %r" % row["pick_rate"]

try:
    billing.preview([lot_id])
    raise AssertionError("счёт по непринятой партии выставлять нельзя")
except ValueError as exc:
    assert "не приняты" in str(exc), exc

# черновик приёмки склад не двигает
supply_applicable = False
import accept

res = accept.legacy_run()
assert res["accepted"] == [], "непроведённая приёмка не должна ставить дату: %r" % res

# приёмка в корзине остаток уже не держит
supply_applicable = True
supply_deleted = "2026-09-02 10:00:00"
res = accept.legacy_run()
assert res["accepted"] == [], "удалённая приёмка не должна ставить дату: %r" % res

supply_deleted = None
res = accept.legacy_run()
assert len(res["accepted"]) == 1, "приёмку не увидел: %r" % res
assert res["accepted"][0]["at"].startswith("2026-09-01T12:30"), res["accepted"][0]

# повторный прогон не переписывает дату и не дёргает МойСклад
before = len(CALLS)
res = accept.legacy_run()
assert res["checked"] == 0 and len(CALLS) == before, "принятые партии больше не опрашиваем"

row = account.lot_rows(only_open=False)[0]
assert row["accepted"] == "01.09.2026", row["accepted"]
assert row["bill_days"] > 0, "после приёмки счётчик должен идти: %r" % row["bill_days"]

# хранение: 10 шт × 2 л × 0,15 ₽ за каждые сутки периода
day = account.storage_of(db.get_lot(lot_id), db.get_client_by_id(client_id), None, account.as_day("2026-09-05"))
assert day["days"] == 4, day
assert day["liter_days"] == 80.0, day
assert day["storage"] == 12.0, day

# сборка: уехало 4 шт по ставке позиции 3 ₽
db.set_lot_pick_rate(lot_id, 3.0)
account.consume_fifo(client_id, PRODUCT_ID, 4, "0001-1", "2026-09-03T10:00:00")
calc = account.money(db.get_lot(lot_id), db.get_client_by_id(client_id), None, account.as_day("2026-09-05"))
assert calc["pick_qty"] == 4.0, calc
assert calc["pick"] == 12.0, calc
assert calc["total"] == round(calc["storage"] + 12.0, 2), calc

prev = billing.preview([lot_id], date_to="2026-09-05")
names = [l["name"] for l in prev["lines"]]
assert names == ["Хранение на складе", "Сборка заказов"], names
assert prev["total"] == round(sum(l["sum"] for l in prev["lines"]), 2), prev

inv = billing.create([lot_id], author="test", date_to="2026-09-05")
payload = [c[2] for c in CALLS if c[1].endswith("/entity/invoiceout")][-1]
assert len(payload["positions"]) == 2, payload["positions"]
assert inv["total"] == prev["total"], (inv, prev)
stored = db.get_invoice(inv["id"])
assert round(stored["ship"], 2) == 12.0, stored["ship"]
assert round(stored["storage"] + stored["ship"], 2) == round(stored["total"], 2), dict(stored)

# после счёта период закрыт: та же партия второй раз те же дни не берёт
again = account.money(db.get_lot(lot_id), db.get_client_by_id(client_id), None, account.as_day("2026-09-05"))
assert again["storage"] == 0 and again["pick"] == 0, again

# выгрузка КиЗ: плоский список, у каждого кода артикул, GTIN и наименование
cab = db.insert_cabinet(client_id, "wb", "wb", "token", "", 1, "", "")
db.replace_cache(client_id, cab, [{
    "marketplace": "wb",
    "ext_key": "1", "ext_article": "ART-1", "ext_barcode": "2000000000019", "name": "Ремень кожаный",
    "size": "", "gtin": "02000000000019", "tracking_type": "Одежда", "subject": "Ремни", "need_kiz": 1, "image": "",
}])
ship_id = db.upsert_shipment(
    client_id, cab, "wb", "fbs", "555", "Новый", "2026-09-04", "ART-1", "2000000000019",
    "Ремень кожаный", 2, None, 2, "2026-09-04T10:00:00", extra={"status_group": "new"},
)
db.replace_shipment_marks(ship_id, [
    {"code": "010200000000001921abcdefgh", "gtin": "02000000000019", "article": "ART-1"},
    {"code": "0104607001112223215xyzqwerty", "gtin": "", "article": ""},
])

name, raw = build_marks([ship_id])
from openpyxl import load_workbook

ws = load_workbook(io.BytesIO(raw)).active
head = [c.value for c in ws[1]]
assert head == ["Код маркировки", "Артикул", "GTIN", "Наименование"], head
first = [c.value for c in ws[2]]
assert first[1:] == ["ART-1", "02000000000019", "Ремень кожаный"], first
second = [c.value for c in ws[3]]
assert second[1] == "ART-1", "артикул подставляем из отправления: %r" % second
assert second[2] == "02000000000019", "GTIN добираем из каталога: %r" % second
assert ws.cell(row=4, column=1).value == "Кодов: 2", ws.cell(row=4, column=1).value

# выгрузки учёта и календаря собираются на новых полях
from export_xlsx import build, build_calendar

xlsx_name, xlsx_raw = build()
head = [c.value for c in load_workbook(io.BytesIO(xlsx_raw)).active[1]]
assert "Ставка хранения, ₽/л/сутки" in head and "Сборка, ₽/шт" in head, head
assert "Приёмка на склад" in head, head
cal_name, cal_raw = build_calendar()
assert cal_raw, "календарь не собрался"

print("хранение, сборка, приёмка и КиЗ: прогон прошёл")

# Регрессии аудита учёта: старые счета, суммы услуг, дата выгрузки, повторы HTTP.
from unittest.mock import patch
from datetime import timedelta
from export_xlsx import build_client_stock, build_invoice

for bad_date in ("не дата", (account.today() + timedelta(days=1)).isoformat()):
    try:
        billing.preview([lot_id], bad_date)
        raise AssertionError("неверная/будущая дата принята")
    except ValueError:
        pass

info = billing.invoice_data(inv["id"])
assert info["invoice"]["pick"] == 12
assert sum(p["total"] for p in info["positions"]) == inv["total"]
original = info["positions"][0]["name"]
conn = db.connect()
conn.execute("UPDATE lots SET name='изменённая карточка',liters=100 WHERE id=?", (lot_id,))
conn.commit(); conn.close()
assert billing.invoice_data(inv["id"])["positions"][0]["name"] == original
assert billing.invoice_data(inv["id"])["positions"][0]["liters"] == 2
report_name, report_raw = build_invoice(inv["id"])
assert load_workbook(io.BytesIO(report_raw)).active.cell(4,9).value == inv["total"]
conn = db.connect(); conn.execute("UPDATE lots SET liters=2 WHERE id=?",(lot_id,)); conn.commit(); conn.close()

# Excel respects the selected cutoff, and never includes another client's rows.
_, raw = build(date_to=account.as_day("2026-09-06"))
ws = load_workbook(io.BytesIO(raw)).active
assert ws.cell(2,20).value == account.lot_rows(date_to=account.as_day("2026-09-06"))[0]["total"]
other = db.insert_client("other", "Другой", "A2", "ORG-1", "STORE-1")
other_lot = db.insert_lot(other, "P2", "=1+1", "2", "", "Другой товар", "", 1, 1, "2026-09-01", "PO2", accepted_at="2026-09-01")
try:
    build_client_stock([lot_id,other_lot])
    raise AssertionError("в клиентский отчёт попал другой клиент")
except ValueError:
    pass
_, raw = build_client_stock([other_lot])
assert load_workbook(io.BytesIO(raw)).active.cell(2,2).data_type == "s"

# Waiting acceptance and future lots cannot have a billed boundary advanced.
waiting_id = db.insert_lot(client_id,"W","W","W","","Ждёт","",1,1,"2026-09-01","PW")
try:
    billing.preview([lot_id,waiting_id], "2026-09-06")
    raise AssertionError("непринятая строка пропущена молча")
except ValueError:
    pass
assert account.row_of(db.get_lot(waiting_id))["days"] == 0
future_id = db.insert_lot(client_id,"F","F","F","","Поздняя","",1,1,"2026-09-10","PF",accepted_at="2026-09-10")
assert len(billing.preview([lot_id,future_id],"2026-09-06")["positions"]) == 1

# Unknown outcome: repeat the exact payload/syncId, even if rate/card changes meanwhile.
real_req = billing.req
attempts = []
def lost_reply(method,url,**kw):
    if method == "POST" and url.endswith("invoiceout"):
        attempts.append(kw["json"])
        raise TimeoutError("ответ потерян")
    return real_req(method,url,**kw)
with patch.object(billing,"req",lost_reply):
    try:
        billing.create([lot_id],date_to="2026-09-06")
        raise AssertionError("ожидался таймаут")
    except TimeoutError:
        pass
assert db.get_lot(lot_id)["billed_until"] == "2026-09-05"
try:
    billing.create([other_lot],date_to="2026-09-06")
    raise AssertionError("незавершённый счёт проигнорирован")
except ValueError:
    pass
def recovered(method,url,**kw):
    if method == "POST" and url.endswith("invoiceout"):
        attempts.append(kw["json"])
        return Fake(200,{"id":"INV-2","name":"СЧ-2"})
    return real_req(method,url,**kw)
with patch.object(billing,"req",recovered):
    result = billing.create([lot_id],date_to="2026-09-06")
assert attempts[0] == attempts[1] and attempts[0]["syncId"]
assert db.get_lot(lot_id)["billed_until"] == "2026-09-06"
assert len(db.list_invoices()) == 2
assert not json.loads(db.get_setting("billing:pending"))

# Atomic rollback after header INSERT; no partial invoice and no period loss.
conn = db.connect()
conn.execute("CREATE TRIGGER test_invoice_fail BEFORE INSERT ON invoice_lots BEGIN SELECT RAISE(ABORT, 'test rollback'); END")
conn.commit();conn.close()
lot = db.get_lot(lot_id); calc = account.row_of(lot,date_to=account.as_day("2026-09-07"))
try:
    db.save_invoice_snapshot(client_id,{"id":"INV-FAIL"},{"storage":1,"intake":0,"ship":0,"total":1},[(lot,calc)],"2026-09-07","2026-09-07","test","2026-09-07")
    raise AssertionError("ожидался отказ записи")
except Exception as exc:
    assert "test rollback" in str(exc), exc
assert len(db.list_invoices()) == 2
assert db.get_lot(lot_id)["billed_until"] == "2026-09-06"

# An old intake-only invoice retains the service and reconciles exactly.
conn=db.connect();conn.execute("DROP TRIGGER test_invoice_fail");conn.commit();conn.close()
old = db.insert_invoice(client_id,"OLD","OLD",0,150,0,150,1,"2026-09-01","test")
conn=db.connect();conn.execute("INSERT INTO invoice_lots(invoice_id,lot_id,intake) VALUES (?,?,150)",(old,lot_id));conn.commit();conn.close()
old_data=billing.invoice_data(old)
assert old_data["positions"][0]["total"] == old_data["invoice"]["total"] == 150
assert old_data["invoice"]["payment_status"] == "Не проверено"
db.set_setting("invoice:ms:%s" % old,json.dumps({"total":150,"paid":50,"applicable":True}))
assert billing.payment_state(db.get_invoice(old))["due"] == 100
assert billing.payment_state(db.get_invoice(old))["payment_status"] == "Частично оплачен"
print("аудит учёта: суммы, снимки, даты, изоляция клиентов, повторы и атомарность — OK")

# Шаблон владельца: числа, оформление, расширение таблицы и реквизиты.
from export_xlsx import build_invoice_document, render_invoice_document
from zipfile import ZipFile
from pathlib import Path
frozen_profile = billing.invoice_profile(inv['id'])
assert frozen_profile['name'] == 'Тест ООО'
_, document = build_invoice_document(inv['id'])
ws = load_workbook(io.BytesIO(document)).active
assert ws['B1'].value == 'БЕРЁЗА ФУЛФИЛМЕНТ'
assert ws['C4'].value == 'Тест ООО'
assert ws['C3'].value == 'с 02.09.2026 по 05.09.2026'
assert [ws.cell(9,c).value for c in range(2,6)] == ['ART-1 Ремень кожаный',3,4,12]
assert ws['E7'].value == stored['storage']
assert ws['E12'].value == inv['total']
assert ws['B23'].value == 'Номер счёта: 00000000000000000001'
assert ws.tables['Table_1'].ref == 'B6:E10'
assert ws.column_dimensions['B'].width == 36.13
assert ws.sheet_view.showGridLines is False
with ZipFile(io.BytesIO(document)) as result_zip, ZipFile(Path(__file__).resolve().parents[1]/'invoice_template.xlsx') as source_zip:
    assert result_zip.read('xl/styles.xml') == source_zip.read('xl/styles.xml')
    assert result_zip.read('xl/theme/theme1.xml') == source_zip.read('xl/theme/theme1.xml')

# Перевыгрузка нового счёта не читает изменившуюся карточку контрагента.
with patch.object(billing, 'client_invoice_profile', side_effect=AssertionError('реквизиты должны быть сохранены')):
    assert build_invoice_document(inv['id'])[1] == document

# Шаблон с большим числом товаров двигает итог и реквизиты, не теряет строки.
profile = {**frozen_profile, 'name': '=1+1'}
for count in (0,1,18,31):
    lines = [{'article':'SKU%s'%n,'name':'Товар %s'%n,'pick_qty':2,'pick_rate':3,'pick':6,'storage':1,'intake':0} for n in range(count)]
    if not lines:
        lines=[{'article':'S','name':'Хранение','pick':0,'storage':4,'intake':0}]
    storage=sum(p['storage'] for p in lines)
    data={'invoice':{'storage':storage,'pick':count*6,'intake':0,'total':storage+count*6,'period_from':'2026-09-01','period_to':'2026-09-05'},'positions':lines}
    _, raw=render_invoice_document(data,profile)
    ws=load_workbook(io.BytesIO(raw)).active
    delta=max(1,count)-18
    assert ws.cell(29+delta,5).value == data['invoice']['total']
    assert ws.cell(40+delta,2).value == 'Номер счёта: 00000000000000000001'
    assert ws['C4'].data_type == 's'
    assert ws.max_row == 45+delta

# Подмена клиента не может взять банковские данные из чужого присланного файла.
db.set_setting('client:invoice-requisites:%s'%client_id,json.dumps({'inn':'другой','account':'ЧУЖОЙ'}))
def no_accounts(method,url,**kw):
    if url.endswith('/accounts'):
        return Fake(200,{'meta':{'size':0},'rows':[]})
    return fake_req(method,url,**kw)
with patch.object(billing,'req',no_accounts):
    assert billing.client_invoice_profile(db.get_client_by_id(client_id))['account'] == ''
# Несколько счетов без основного: не выбираем случайный.
def ambiguous_accounts(method,url,**kw):
    if url.endswith('/accounts'):
        return Fake(200,{'meta':{'size':2},'rows':[{'accountNumber':'1'},{'accountNumber':'2'}]})
    return fake_req(method,url,**kw)
with patch.object(billing,'req',ambiguous_accounts):
    try:
        billing.client_invoice_profile(db.get_client_by_id(client_id))
        raise AssertionError('выбран случайный банковский счёт')
    except ValueError:
        pass

# Ошибка подготовки Excel не создаёт финансовый документ и не закрывает период.
before_count=len(db.list_invoices())
with patch('export_xlsx.render_invoice_document',side_effect=ValueError('шаблон недоступен')):
    try:
        billing.create([lot_id],date_to='2026-09-07')
        raise AssertionError('ожидался отказ до МойСклад')
    except ValueError as exc:
        assert 'шаблон недоступен' in str(exc)
assert len(db.list_invoices()) == before_count
assert db.get_lot(lot_id)['billed_until'] == '2026-09-06'
# Объединение артикулов меняет представление, сохраняя все исходные партии и ставки.
base = {'client_id': 1, 'article': 'ONE', 'name': 'Товар', 'ms_product_id': 'P',
        'liters': 2, 'pick_rate': 34, 'qty': 0, 'qty_in': 0, 'storage': 0,
        'pick': 0, 'pick_qty': 0, 'total': 0, 'billing_errors': []}
raw_rows = [
    {**base, 'id': 101, 'kind': 'receipt', 'qty_in': 10, 'qty': 5, 'storage': 8.25,
     'total': 8.25, 'bill_days': 7, 'bill_from': '2026-10-05', 'bill_to': '2026-10-11'},
    {**base, 'id': 102, 'kind': 'assembly', 'pick_qty': 2, 'pick': 68, 'total': 68},
    {**base, 'id': 103, 'kind': 'assembly', 'pick_qty': 3, 'pick': 102, 'total': 102},
]
grouped = account.group_rows(raw_rows)
assert len(grouped) == 1 and grouped[0]['lot_ids'] == [101, 102, 103]
assert grouped[0]['pick_qty'] == 5 and grouped[0]['total'] == 178.25
assert grouped[0]['bill_days'] == 7
assert account.group_rows(grouped) == grouped
assert len(account.group_rows(raw_rows + [{**base, 'id': 104, 'client_id': 2}])) == 2
assert account.group_rows(raw_rows + [{**base, 'id': 105, 'billing_errors': ['не заполнен литраж']}])[0]['billing_errors']
group_data = {'invoice': {'storage': 8.25, 'pick': 170, 'intake': 0, 'total': 178.25,
                          'period_from': '2026-10-05', 'period_to': '2026-10-11'}, 'positions': raw_rows}
_, raw = render_invoice_document(group_data, frozen_profile)
ws = load_workbook(io.BytesIO(raw)).active
assert [ws.cell(9,c).value for c in range(2,6)] == ['ONE Товар', 34, 5, 170]
assert sum(str(cell.value or '').startswith('ONE ') for row in ws for cell in row) == 1
# Разные реальные ставки сохраняются в одной строке без выдуманной средней цены.
raw_rows[2].update(pick_rate=35, pick=105, total=105)
group_data['invoice'].update(pick=173, total=181.25)
_, raw = render_invoice_document(group_data, frozen_profile)
ws = load_workbook(io.BytesIO(raw)).active
assert [ws.cell(9,c).value for c in range(3,6)] == ['По ставкам', 5, 173]
assert '2 шт × 34 ₽' in ws['B9'].value and '3 шт × 35 ₽' in ws['B9'].value
raw_rows[2]['pick_qty'] = None
try:
    render_invoice_document(group_data, frozen_profile)
    raise AssertionError('неизвестное количество нельзя заменять нулём')
except ValueError:
    pass
print('счёт по шаблону: суммы, стили, объединение артикулов, разные ставки, реквизиты и повторы — OK')

# Добавление контрагента: дубли всего аккаунта, секреты, ошибка ключа и повтор.
import copy
import agents_sync as onboarding
os.environ['MS_ORG_ID']='ORG-1'
os.environ['MS_STORE_ID']='STORE-1'
for _name,_typ,key in onboarding.AGENT_ATTRS:
    db.set_setting(key,'META-'+key)
entities={}
mutations=[]
reject_wb=False
lose_create=False

def onboarding_req(method,url,headers=None,**kw):
    global lose_create
    if 'suggestions.dadata.ru' in url:
        return Fake(200,{'suggestions':[{'data':{'inn':'7730588444','kpp':'770501001','ogrn':'1087746982157','type':'LEGAL','state':{'status':'ACTIVE'},'name':{'short_with_opf':'ООО Пример','full_with_opf':'Общество Пример'},'address':{'unrestricted_value':'Адрес из справочника'}}}]})
    if url.endswith('/ping'):
        return Fake(401 if reject_wb else 200,{})
    if '/content/v2/get/cards/list' in url:
        return Fake(200,{'cards':[]})
    if '/api/v3/warehouses' in url:
        return Fake(200,[])
    if '/v3/product/list' in url:
        return Fake(200,{'result':{'items':[]}})
    root=net.MS_BASE+'/entity/counterparty'
    if url==root and method=='GET':
        filt=kw['params']['filter']; inn=filt.split(';')[0].split('=')[1]
        archived='archived=true' in filt
        rows=[copy.deepcopy(e) for e in entities.values() if e.get('inn')==inn and bool(e.get('archived'))==archived]
        return Fake(200,{'meta':{'size':len(rows)},'rows':rows})
    if method in ('POST','PUT') and url.startswith(root):
        data=copy.deepcopy(kw['json']);mutations.append((method,data))
        eid=url.rsplit('/',1)[-1] if method=='PUT' else 'NEW-%s'%(len(entities)+1)
        entity=entities.setdefault(eid,{'id':eid})
        attrs={a['meta']['href']:a for a in entity.get('attributes',[])}
        for a in data.get('attributes',[]):
            a['name']=next(n for n,t,k in onboarding.AGENT_ATTRS if a['meta']['href'].endswith('META-'+k))
            attrs[a['meta']['href']]=a
        entity.update({k:v for k,v in data.items() if k!='attributes'})
        entity['attributes']=list(attrs.values())
        if method=='POST' and lose_create:
            lose_create=False
            raise TimeoutError('ответ потерян')
        return Fake(200,copy.deepcopy(entity))
    if method=='GET' and url.startswith(root+'/'):
        return Fake(200,copy.deepcopy(entities[url.rsplit('/',1)[-1]]))
    raise AssertionError((method,url))

form={'inn':'7730588444','companyType':'legal','name':'ООО Пример','legalTitle':'Общество Пример','legalAddress':'Адрес компании',
      'kpp':'770501001','ogrn':'1087746982157','tariff_pick':'12,50','wb_token':'verified-wb-secret','ozon_client_id':'123456','ozon_key':'verified-ozon-secret'}
with patch.object(onboarding,'req',onboarding_req):
    for bad in ('0000000000','7730588445','123','1;archived=true'):
        try:onboarding.checked_inn(bad);raise AssertionError('принят неверный ИНН')
        except ValueError:pass
    assert onboarding.checked_inn('7730588444') == '7730588444'
    lookup=onboarding.lookup_company(form['inn'])
    assert lookup['existing']==[] and lookup['company'] is None
    with patch.dict(os.environ,{'DADATA_API_KEY':'test-private-key'}):
        lookup=onboarding.lookup_company(form['inn'])
        assert lookup['company']['legalAddress']=='Адрес из справочника'
    # Неверный ключ не создаёт компанию и не меняет локальную базу.
    before=len(db.list_clients());reject_wb=True
    try:onboarding.onboard_client(form);raise AssertionError('неверный WB токен принят')
    except ValueError as exc:assert 'WB' in str(exc)
    assert not mutations and len(db.list_clients())==before
    reject_wb=False
    made=onboarding.onboard_client(form)
    assert made['created'] and made['connected']==['WB','Ozon']
    local=db.get_client_by_id(made['id'])
    assert local['tariff_pick']==12.5
    assert db.get_cabinet_by_client_mp(local['id'],'wb')['token']==form['wb_token']
    assert db.get_cabinet_by_client_mp(local['id'],'ozon')['client_id_ext']=='123456'
    assert form['wb_token'] not in json.dumps(mutations,ensure_ascii=False)
    assert form['ozon_key'] not in json.dumps(mutations,ensure_ascii=False)
    duplicate=onboarding.lookup_company(form['inn'])
    assert duplicate['existing'][0]['client_id']==made['id']
    assert 'token' not in json.dumps(duplicate) and form['wb_token'] not in json.dumps(duplicate)
    # Повтор без выбранной карточки показывает совпадение, ничего не создаёт.
    before=len(mutations)
    assert onboarding.onboard_client(form)['existing']
    assert len(mutations)==before
    old_id=duplicate['existing'][0]['id']
    reuse={**form,'existing_id':old_id,'name':'Не подменять имя','wb_token':'','ozon_client_id':'','ozon_key':'','tariff_pick':''}
    reused=onboarding.onboard_client(reuse)
    assert reused['id']==made['id'] and not reused['created']
    assert entities[old_id]['name']==form['name']
    assert db.get_cabinet_by_client_mp(made['id'],'wb')['token']==form['wb_token']
    assert db.get_client_by_id(made['id'])['tariff_pick']==12.5
    # Архивный дубль нельзя обойти созданием нового.
    entities['ARCHIVED']={**copy.deepcopy(entities[old_id]),'id':'ARCHIVED','archived':True}
    assert len(onboarding.lookup_company(form['inn'])['existing'])==2
    try:onboarding.onboard_client({**reuse,'existing_id':'ARCHIVED'});raise AssertionError('архивная карточка принята')
    except ValueError as exc:assert 'архив' in str(exc)
    # Другому клиенту нельзя передать тот же кабинет.
    other_form={**form,'inn':'7707083893','kpp':'773601001','ogrn':'1027700132195'}
    before=len(mutations)
    try:onboarding.onboard_client(other_form);raise AssertionError('кабинет присвоен другому клиенту')
    except ValueError as exc:assert 'другому клиенту' in str(exc)
    assert len(mutations)==before
    # Ответ потерян после фактического создания: новый дубль не появляется.
    lose_create=True
    lost={**other_form,'wb_token':'','ozon_client_id':'','ozon_key':''}
    try:onboarding.onboard_client(lost);raise AssertionError('ожидалась потеря ответа')
    except TimeoutError:pass
    assert len([e for e in entities.values() if e['inn']==lost['inn']])==1
    retry=onboarding.onboard_client(lost)
    assert len(retry['existing'])==1
    recovered=onboarding.onboard_client({**lost,'existing_id':retry['existing'][0]['id']})
    assert not recovered['created']
    assert len([e for e in entities.values() if e['inn']==lost['inn']])==1
    # Любая локальная ошибка откатывает и клиента, и кабинеты.
    conn=db.connect();conn.execute("CREATE TRIGGER fail_onboard BEFORE INSERT ON cabinets BEGIN SELECT RAISE(ABORT,'atomic onboarding'); END");conn.commit();conn.close()
    before=len(db.list_clients())
    try:
        db.save_onboard_client({'id':'LOCAL-FAIL','name':'Локальный тест'},'LOCALFAIL','ORG-1','STORE-1',{'wb_token':'private-test','ozon_key':'','ozon_client_id':''},1)
        raise AssertionError('ожидался отказ записи')
    except Exception as exc:assert 'atomic onboarding' in str(exc)
    assert len(db.list_clients())==before
    conn=db.connect();conn.execute('DROP TRIGGER fail_onboard');conn.commit();conn.close()
print('контрагенты: ИНН, справочник, дубли, архив, ключи, потеря ответа и атомарность — OK')

# Банковские реквизиты: необязательны, но частичный ввод не принимается.
assert onboarding._onboard_bank({}) is None
try:onboarding._onboard_bank({'bic':'123456789'});raise AssertionError('частичный банк принят')
except ValueError:pass
bank={'accountNumber':'40802810000000000001','bic':'044525225','bankName':'Банк для проверки','correspondentAccount':'30101810000000000002'}
assert onboarding._onboard_bank(bank)==bank
bank_rows=[];bank_posts=[];bank_lost=True

def bank_req(method,path,**kw):
    global bank_lost
    if method=='GET':return {'meta':{'size':len(bank_rows)},'rows':copy.deepcopy(bank_rows)}
    assert method=='POST' and kw['retry_safe'] is False
    bank_posts.append(kw['json']);bank_rows.append(copy.deepcopy(kw['json']))
    if bank_lost:
        bank_lost=False
        raise TimeoutError('потерян ответ банка')
    return kw['json']
with patch.object(onboarding,'_agent_request',bank_req):
    try:onboarding._save_onboard_bank('test',bank);raise AssertionError('ожидался обрыв')
    except TimeoutError:pass
    onboarding._save_onboard_bank('test',bank)
    assert len(bank_posts)==1 and len(bank_rows)==1
    try:onboarding._save_onboard_bank('test',{**bank,'accountNumber':'40802810000000000003'});raise AssertionError('заменён банк')
    except ValueError:pass
assert onboarding.checked_inn('770300584079')=='770300584079'
print('банковские реквизиты: частичный ввод, сохранение и повтор после обрыва — OK')

# Пустая цена сборки не означает бесплатную услугу. Явный ноль допустим.
rate_client = db.insert_client('rate-check', 'Проверка ставки', 'AGENT-RATE', 'ORG-1', 'STORE-1')
rate_lot = db.insert_lot(rate_client, 'RATE-PRODUCT', 'RATE', '', '', 'Проверка', '', 1, 3,
                         '2026-09-01T10:00:00', '', accepted_at='2026-09-01T10:00:00')
db.add_lot_move(rate_lot, 1, 'ship', 'RATE-DEMAND', '2026-09-02T10:00:00')
rate_row = account.row_of(db.get_lot(rate_lot), date_to=account.as_day('2026-09-03'))
assert rate_row['billing_errors'] == ['не заполнена ставка сборки']
assert account.totals([rate_row])['incomplete'] == 1
for action in (lambda: billing.preview([rate_lot], '2026-09-03'),
               lambda: __import__('export_xlsx').build_client_stock([rate_lot], account.as_day('2026-09-03'))):
    try:
        action()
        raise AssertionError('отсутствующая ставка скрыто принята за ноль')
    except ValueError as exc:
        assert 'ставк' in str(exc)
db.set_lot_pick_rate(rate_lot, 0)
rate_row = account.row_of(db.get_lot(rate_lot), date_to=account.as_day('2026-09-03'))
assert rate_row['billing_errors'] == [] and rate_row['pick_rate_known']
assert billing.preview([rate_lot], '2026-09-03')['pick_qty'] == 1
print('сборка: отсутствующая ставка блокирует счёт и отчёт; явный ноль допустим — OK')

# Полный складской источник: отдельные счета при общем товаре группы.
os.environ['FF_DB'] = os.path.join(tempfile.mkdtemp(), 'ledger.db')
db.init_db()
aid=db.insert_client('A','Приёмка','A','ORG','S');bid=db.insert_client('B','Отгрузка','B','ORG','S')
def ledger_fixture(qty=4, same_day=False):
    return {'products':{'P':{'id':'P','name':'Товар','article':'P','barcode':'','gtin':'','tracking_type':'NOT_TRACKED','liters':2,'pick_rate':3}},
            'parties':{'A':{'id':'A','name':'Приёмка'},'B':{'id':'B','name':'Отгрузка'}},
            'documents':[{'id':'IN','kind':'supply','name':'IN','moment':'2026-09-10T10:00:00','party':'A','store':'S','positions':[{'id':'I1','product':'P','qty':10}]},
                         {'id':'OUT','kind':'demand','name':'OUT','moment':'2026-09-%sT12:00:00'%('10' if same_day else '12'),'party':'B','store':'S','positions':[{'id':'O1','product':'P','qty':qty}]}],
            'stock':{'P':10-qty}}
snap=ledger_fixture();groups=[['A','B']];plan=accept.build_plan(snap,groups)
accept.apply_plan(snap,plan,groups)
receipt=next(x for x in db.list_lots() if x['ledger_kind']=='receipt');assembly=next(x for x in db.list_lots() if x['ledger_kind']=='assembly')
assert receipt['client_id']==aid and assembly['client_id']==bid
r=account.row_of(receipt,date_to=account.as_day('2026-09-12'));a=account.row_of(assembly,date_to=account.as_day('2026-09-12'))
assert r['bill_days']==2 and r['liter_days']==40 and r['storage']==6 and r['pick']==0,r
assert a['storage']==0 and a['pick_qty']==4 and a['pick']==12,a
assert billing.preview([receipt['id']],'2026-09-12')['total']==6
assert billing.preview([assembly['id']],'2026-09-12')['total']==12
cal=account.calendar(date_from=account.as_day('2026-09-10'),date_to=account.as_day('2026-09-12'))
assert len(cal['rows'])==1 and [c['sum'] for c in cal['rows'][0]['cells']]==[0,3,3],cal
assert cal['total']==6
ids=[x['id'] for x in db.list_lots()];conn=db.connect();n_moves=conn.execute('select count(*) from lot_moves').fetchone()[0];conn.close()
accept.apply_plan(snap,accept.build_plan(snap,groups),groups)
assert [x['id'] for x in db.list_lots()]==ids
conn=db.connect();assert conn.execute('select count(*) from lot_moves').fetchone()[0]==n_moves;assert conn.execute('select count(*) from ms_ledger_history').fetchone()[0]==1;conn.close()
try:accept.build_plan(snap,[]);raise AssertionError('смешаны чужие группы')
except ValueError as e:assert 'не хватает' in str(e)
# День приёмки бесплатный, сборка в этот день платная.
same=ledger_fixture(qty=10,same_day=True);accept.apply_plan(same,accept.build_plan(same,groups),groups)
rr=account.row_of(db.get_lot(receipt['id']),date_to=account.as_day('2026-09-12'))
aa=account.row_of(db.get_lot(assembly['id']),date_to=account.as_day('2026-09-12'))
assert rr['storage']==0 and rr['bill_days']==0 and aa['pick']==30,(rr,aa)
# Отмена/распроведение источника восстанавливает остаток и архивирует начисление.
cancel=ledger_fixture();cancel['documents']=cancel['documents'][:1];cancel['stock']={'P':10}
accept.apply_plan(cancel,accept.build_plan(cancel,groups),groups)
assert len(db.list_lots())==1 and db.list_lots()[0]['qty_left']==10
assert db.get_lot(assembly['id'])['source_active']==0
# Повторные строки приёмки не схлопываются по артикулу.
dup=ledger_fixture();dup['documents'][0]['positions']=[{'id':'I1','product':'P','qty':5},{'id':'I2','product':'P','qty':5}]
p=accept.build_plan(dup,groups);assert len([x for x in p['lots'].values() if x['ledger_kind']=='receipt'])==2
# Тариф клиента не подменяет отсутствующее допполе товара.
missing=ledger_fixture();missing['products']['P']['pick_rate']=None;missing['products']['P']['liters']=None
accept.apply_plan(missing,accept.build_plan(missing,groups),groups);db.update_client(bid,tariff_pick=999)
try:billing.preview([assembly['id']],'2026-09-12');raise AssertionError('подставлен тариф клиента')
except ValueError as e:assert 'ставк' in str(e)
# Сборку можно выставлять без литража, когда её цена заполнена.
missing['products']['P']['pick_rate']=3
accept.apply_plan(missing,accept.build_plan(missing,groups),groups)
assert billing.preview([assembly['id']],'2026-09-12')['total']==12
try:billing.preview([receipt['id']],'2026-09-12');raise AssertionError('пустой литраж принят за бесплатное хранение')
except ValueError as e:assert 'литраж' in str(e)
assert account.calendar(date_from=account.as_day('2026-09-10'),date_to=account.as_day('2026-09-12'))['total'] is None
# Изменение закрытого периода не переписывает учёт и границы счёта.
accept.apply_plan(snap,accept.build_plan(snap,groups),groups)
db.mark_lot_billed(assembly['id'],'2026-09-12',0,0,4)
changed=ledger_fixture(qty=5)
try:accept.apply_plan(changed,accept.build_plan(changed,groups),groups);raise AssertionError('изменён выставленный период')
except ValueError as e:assert 'закрытого периода' in str(e)
assert db.get_lot(assembly['id'])['billed_until']=='2026-09-12'
assert account.row_of(db.get_lot(receipt['id']),date_to=account.as_day('2026-09-12'))['qty']==6
assert billing.preview([assembly['id']],'2026-09-12')['total']==0
# Ошибка/неполная страница не выдаётся за успешный пустой ответ.
with patch.object(accept,'req',lambda *a,**k: Fake(200,{'meta':{'size':2},'rows':[]})):
    try:accept.ms_pages('/bad');raise AssertionError('принята неполная страница')
    except ValueError:pass
assert db.connect().execute('pragma quick_check').fetchone()[0]=='ok'
print('МойСклад: группы, раздельные начисления, даты, отмены, дубли строк, пустые поля, повтор и закрытые периоды — OK')

# Цена из «Заведения» записывается в карточку; общая карточка не меняет владельца.
import products_push as pp
import ms
from copy import deepcopy
db.set_setting('ledger:groups',json.dumps([['A','B']]))
db.set_setting('ATTR_PICK_RATE_ID','COST');db.set_setting('ATTR_LITERS_ID','LIT');db.set_setting('ATTR_CLIENT_ID','OWNER')
shared={'id':'SHARED','article':'P','name':'Общий товар','supplier':ms.ms_meta('counterparty','A'),
        'attributes':[{'name':'Клиент фулфилмента','meta':{'href':net.MS_BASE+'/entity/product/metadata/attributes/OWNER'},'value':ms.ms_meta('counterparty','A')}]}
assert pp.same_client_product(shared,db.get_client_by_id(bid))
try:pp.same_client_product(shared,{'ms_counterparty_id':'OUTSIDER'});raise AssertionError('разрешена чужая группа')
except ValueError:pass
sent=[]
def product_saved(method,url,headers=None,**kw):
    sent.append((method,deepcopy(kw['json'])))
    return Fake(200,{'id':'SHARED','attributes':kw['json']['attributes']})
with patch.object(pp,'get_cabinet',return_value={'id':1,'client_id':bid,'marketplace':'wb'}),patch.object(pp,'find_linked',return_value=deepcopy(shared)),patch.object(pp,'link_hits'),patch.object(pp,'req',product_saved):
    result=pp.push_one({'cabinet_id':1,'name':'Общий товар','Литраж_л':2,'Стоимость сборки':7})
    assert result['ok'],result
assert sent[0][0]=='PUT' and 'supplier' not in sent[0][1]
attrs={pp._ref_id(x['meta']):x['value'] for x in sent[0][1]['attributes']}
assert attrs['COST']==7 and attrs['LIT']==2 and 'OWNER' not in attrs
zero=ms.product_attrs(2,'B',0);assert any(pp._ref_id(x['meta'])=='COST' and x['value']==0 for x in zero)
print('Заведение: стоимость сборки сохраняется в МойСклад; карточка общей компании сохраняет владельца — OK')

# Редактирование из учёта: точный UUID, частичная запись, проверка ответа и пересчёт.
pid = '00000000-0000-0000-0000-000000000001'
conn = db.connect()
conn.execute('UPDATE lots SET ms_product_id=?', (pid,))
conn.commit(); conn.close()
edit_lot = db.list_lots(only_open=False)[0]['id']
remote = {'id':pid, 'name':'Общий товар', 'article':'P', 'updated':'v1', 'attributes':[
    {'name':'Литраж, л','meta':{'href':net.MS_BASE+'/entity/product/metadata/attributes/LIT'},'value':2},
    {'name':'Стоимость сборки, ₽','meta':{'href':net.MS_BASE+'/entity/product/metadata/attributes/COST'},'value':3},
    {'name':'Клиент фулфилмента','meta':{'href':net.MS_BASE+'/entity/product/metadata/attributes/OWNER'},'value':ms.ms_meta('counterparty','A')}]}
edit_writes=[]
lose_response=False
def edit_get(path):
    if path.endswith('/metadata/attributes'):
        return {'rows':[{'id':'LIT','name':'Литраж, л','type':'double'}, {'id':'COST','name':'Стоимость сборки, ₽','type':'double'}]}
    assert path == '/entity/product/' + pid
    return deepcopy(remote)
def edit_put(method,url,**kw):
    assert method=='PUT' and url.endswith('/'+pid) and kw['retry_safe'] is False
    assert set(kw['json'])=={'attributes'}
    edit_writes.append(deepcopy(kw['json']))
    for sent_attr in kw['json']['attributes']:
        aid=accept.ref_id(sent_attr)
        existing=next(a for a in remote['attributes'] if accept.ref_id(a)==aid)
        existing['value']=sent_attr['value']
    remote['updated']+='x'
    if lose_response:raise OSError('потерян ответ')
    return Fake(200,deepcopy(remote))
def edit_snapshot():
    return {'products':{pid:{'liters':next(a['value'] for a in remote['attributes'] if a['name']=='Литраж, л'),
                             'pick_rate':next(a['value'] for a in remote['attributes'] if a['name']=='Стоимость сборки, ₽')}}}
db.set_setting('billing:pending','null')
with patch.object(accept,'ms_get',edit_get),patch.object(net,'req',edit_put),patch.object(accept,'fetch_snapshot',edit_snapshot),patch.object(accept,'build_plan',return_value={}),patch.object(accept,'apply_plan',return_value={'last_sync':'now'}) as applied:
    card=account.lot_product(edit_lot)
    assert card['liters']==2 and card['pick_rate']==3
    for fields in ({'liters':'0'}, {'liters':'-1'}, {'liters':'NaN'}, {'liters':'inf'}, {'pick_rate':'-1'}, {'pick_rate':'nan'}, {'pick_rate':True}, {}):
        try:account.save_lot_product(edit_lot,dict(fields,revision=card['revision']),'tester');raise AssertionError(fields)
        except ValueError:pass
    assert not edit_writes
    try:account.save_lot_product(edit_lot,{'liters':'5','revision':'stale'},'tester');raise AssertionError('принята устаревшая карточка')
    except ValueError:pass
    db.set_setting('billing:pending','{"syncId":"pending"}')
    try:account.save_lot_product(edit_lot,{'liters':'5','revision':card['revision']},'tester');raise AssertionError('изменение при неизвестном счёте')
    except ValueError:pass
    db.set_setting('billing:pending','null')
    result=account.save_lot_product(edit_lot,{'liters':'8х8х120','pick_rate':'','revision':card['revision']},'tester')
    assert result['synced'] and result['product']['liters']==7.68 and result['product']['pick_rate']==3
    assert len(edit_writes[-1]['attributes'])==1 and accept.ref_id(edit_writes[-1]['attributes'][0])=='LIT'
    assert next(a for a in remote['attributes'] if a['name']=='Клиент фулфилмента')['value']==ms.ms_meta('counterparty','A')
    lose_response=True
    result=account.save_lot_product(edit_lot,{'pick_rate':'0','revision':result['product']['revision']},'tester')
    assert result['synced'] and result['product']['pick_rate']==0
    assert applied.call_count==2
    with patch.object(accept,'fetch_snapshot',side_effect=ValueError('неполный источник')):
        result=account.save_lot_product(edit_lot,{'pick_rate':'4,50','revision':result['product']['revision']},'tester')
        assert result['saved'] and not result['synced'] and 'неполный источник' in db.get_setting('ledger:last_error')
        assert applied.call_count==2
    try:account._product_for_lot(-1);raise AssertionError('допущена неизвестная партия')
    except ValueError:pass
print('Учёт: запись допполей, габариты, нулевая сборка, защита от конфликта/счёта, потеря ответа и неполная сверка — OK')

# Финансовые endpoints требуют явную дату; пустой фильтр не становится сегодняшним днём.
import web
from fastapi import HTTPException
with patch.object(web, 'who', return_value='tester'), patch.object(web, 'init_db'):
    for endpoint, function in ((web.invoice_preview, 'preview'), (web.invoice_create, 'create')):
        with patch.object(billing, function, return_value={'period_to':'2026-10-11'}) as action:
            for value in ('', None, 'не дата', '2099-01-01'):
                try:
                    endpoint({'lot_ids':[lot_id], 'date_to':value}, '')
                    raise AssertionError('неявная или неверная дата допустила счёт')
                except HTTPException as exc:
                    assert exc.status_code == 400
            action.assert_not_called()
            result = endpoint({'lot_ids':[lot_id], 'date_to':'2026-10-11'}, '')
            assert result['period_to'] == '2026-10-11'
            assert action.call_args.args[-1] == '2026-10-11'
print('Счёт: обязательная дата, неверная дата, будущий период и точная передача выбранного окончания — OK')
