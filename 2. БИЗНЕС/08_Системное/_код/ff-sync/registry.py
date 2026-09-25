"""Реестр входящей поставки от клиента: xlsx или docx → шапка поставки и позиции по артикулу.

Формат живой: клиенты присылают таблицу с разным числом пустых колонок и своими
формулировками в шапке. Поэтому колонки ищем по словам в строке заголовка, а не
по номеру, а шапку собираем из всех строк вне таблицы позиций.
"""

import io
import re

from db import barcode_norm

# Слово в заголовке колонки → ключ позиции.
COLUMNS = (
    ("артикул", "article"),
    ("штрихкод", "barcode"),
    ("штрих-код", "barcode"),
    ("штрих код", "barcode"),
    ("наименование", "name"),
    ("название", "name"),
    ("кол-во", "qty"),
    ("количество", "qty"),
    ("мест", "places"),
    ("ед.", "unit"),
    ("ед,", "unit"),
    ("единиц", "unit"),
    ("примечан", "note"),
)

# Слово в первой ячейке строки вне таблицы → поле шапки поставки.
# Порядок значим: «Контакт клиента» должен попасть в contact, а не в client_name.
HEAD = (
    ("№ поставки", "number"),
    ("номер поставки", "number"),
    ("контакт", "contact"),
    ("клиент", "client_name"),
    ("организац", "client_name"),
    ("договор", "contract"),
    ("плановая дата", "planned_at"),
    ("дата и время", "planned_at"),
    ("перевозчик", "carrier"),
    ("водител", "carrier"),
    ("гос. номер", "car_plate"),
    ("гос номер", "car_plate"),
    ("грузовых мест", "places"),
    ("маркировка", "marking"),
    ("честный знак", "marking"),
    ("особые условия", "terms"),
)


def norm(raw):
    text = str(raw if raw is not None else "").replace("\u00a0", " ").replace("ё", "е")
    return re.sub(r"\s+", " ", text).strip().lower()


def cells(row):
    return ["" if v is None else str(v).replace("\u00a0", " ").strip() for v in row]


BLANK = re.compile(r"_{2,}")
# Подписи пустого бланка. «ФИО ____ тел. ____» — это не контакт.
FORM = {"фио", "тел", "тел.", "дата", "должность"}


def clean_value(raw):
    """Убирает незаполненные прочерки бланка: «Паллет: 1  Коробов: ___» → «Паллет: 1»."""
    text = BLANK.sub("", str(raw or ""))
    parts = [p.strip() for p in re.split(r"\s{2,}", text)]
    keep = [p for p in parts if p and not p.endswith(":") and norm(p) not in FORM]
    return re.sub(r"\s+", " ", " ".join(keep)).strip()


def number(raw):
    text = str(raw if raw is not None else "").strip().replace("\u00a0", "").replace(" ", "")
    text = text.replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def is_header(row):
    joined = [norm(c) for c in row]
    has_article = any("артикул" in c for c in joined)
    has_code = any("штрихкод" in c or "штрих-код" in c or "штрих код" in c for c in joined)
    has_name = any("наименование" in c or "название" in c for c in joined)
    return has_article and (has_code or has_name)


def map_columns(row):
    """Заголовок таблицы → {ключ: индекс колонки}. Первое совпадение выигрывает."""
    found = {}
    for idx, raw in enumerate(row):
        text = norm(raw)
        if not text:
            continue
        for word, key in COLUMNS:
            if word in text and key not in found:
                found[key] = idx
                break
    return found


def head_field(label):
    text = norm(label)
    if not text:
        return ""
    for word, key in HEAD:
        if word in text:
            return key
    return ""


def head_pairs(row):
    """Строка шапки может нести одну пару или две: подпись, значение, подпись, значение.

    Следующая подпись не считается значением предыдущего поля.
    """
    filled = [v for v in cells(row) if v]
    out = []
    i = 0
    while i < len(filled):
        key = head_field(filled[i])
        if not key:
            i += 1
            continue
        if i + 1 < len(filled) and not head_field(filled[i + 1]):
            val = clean_value(filled[i + 1])
            i += 2
        else:
            val = ""
            i += 1
        if val:
            if key == "planned_at":
                val = tidy_date(val)
            out.append((key, val))
    return out


def tidy_date(val):
    hit = re.match(r"^(\d{1,2})\s*[./]\s*(\d{1,2})\s*[./]\s*(\d{4})$", val.strip())
    if not hit:
        return val
    day, month, year = hit.groups()
    return "%02d.%02d.%s" % (int(day), int(month), year)


W_ML = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def read_docx(data):
    import zipfile
    from xml.etree import ElementTree as ET

    try:
        xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")
    except Exception as exc:
        raise ValueError("не смог прочитать файл Word") from exc
    root = ET.fromstring(xml)
    rows = []
    for tbl in root.iter(W_ML + "tbl"):
        for tr in tbl.findall(W_ML + "tr"):
            line = []
            for tc in tr.findall(W_ML + "tc"):
                line.append("".join(t.text or "" for t in tc.iter(W_ML + "t")).strip())
            rows.append(line)
    return rows


def read_rows(filename, data):
    name = (filename or "").lower()
    if name.endswith(".docx"):
        return read_docx(data)
    if not (name.endswith(".xlsx") or name.endswith(".xlsm") or name.endswith(".xls")):
        raise ValueError("реестр ждём файлом Excel или Word: xlsx, xls или docx")
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows(values_only=True):
            out.append(list(row))
        if any(is_header(r) for r in out):
            break
        out = []
    return out


def position(row, cols):
    def cell(key):
        idx = cols.get(key)
        if idx is None or idx >= len(row):
            return ""
        return row[idx]

    article = str(cell("article") or "").strip()
    barcode = str(cell("barcode") or "").strip()
    if barcode.endswith(".0") and barcode[:-2].isdigit():
        barcode = barcode[:-2]
    if not article and not barcode:
        return None
    return {
        "article": article,
        "barcode": barcode,
        "name": str(cell("name") or "").strip(),
        "unit": str(cell("unit") or "").strip(),
        "qty": number(cell("qty")) or 0.0,
        "places": number(cell("places")),
        "note": str(cell("note") or "").strip(),
    }


def merge_same(rows):
    """Один артикул двумя строками — это две коробки одного товара, а не два товара."""
    order = []
    seen = {}
    merged = 0
    for item in rows:
        key = barcode_norm(item["barcode"]) or item["article"].upper()
        if key in seen:
            first = seen[key]
            first["qty"] += item["qty"]
            if item["places"] is not None:
                first["places"] = (first["places"] or 0) + item["places"]
            merged += 1
            continue
        seen[key] = item
        order.append(item)
    return order, merged


def parse(filename, data):
    rows = read_rows(filename, data)
    start = next((i for i, row in enumerate(rows) if is_header(row)), -1)
    if start < 0:
        raise ValueError("не нашёл шапку таблицы: нужны колонки «Артикул» и «Штрихкод»")
    cols = map_columns(rows[start])
    if "qty" not in cols:
        raise ValueError("не нашёл колонку с количеством")

    positions = []
    stop = len(rows)
    for i in range(start + 1, len(rows)):
        item = position(rows[i], cols)
        if item is None:
            if positions:
                stop = i
                break
            continue
        positions.append(item)

    supply = {}
    for i, row in enumerate(rows):
        if start <= i < stop:
            continue
        for key, val in head_pairs(row):
            if key and val and not supply.get(key):
                supply[key] = val

    positions, merged = merge_same(positions)
    problems = []
    for item in positions:
        if item["qty"] <= 0:
            problems.append("%s: нет количества" % (item["article"] or item["barcode"]))
    if not positions:
        raise ValueError("в реестре нет ни одной позиции")
    return {
        "supply": supply,
        "rows": positions,
        "merged": merged,
        "problems": problems,
        "qty_total": sum(item["qty"] for item in positions),
    }
