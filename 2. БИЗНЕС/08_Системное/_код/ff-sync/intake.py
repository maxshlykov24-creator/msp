"""Приёмка: список штрихкодов → сверка с кабинетами → товары в МойСклад."""

import re

import registry
from db import (
    QUEUE_STATES,
    add_intake_row,
    barcode_norm,
    delete_supply,
    find_cache,
    find_cache_group,
    get_client_by_id,
    get_intake_row,
    init_db,
    insert_lot,
    insert_supply,
    list_cabinets,
    list_intake,
    prefer_hit,
    run_lock,
    update_intake_row,
    upsert_sku,
)
from ms import kind_from_subject, tracking_code
from products_push import create_purchase_order, gtin14, push_one

from datetime import datetime, timedelta, timezone

MSK = timezone(timedelta(hours=3))
BARCODE_RE = re.compile(r"(OZN[A-Z0-9]+|\d{8,14})", re.I)


def now_iso():
    return datetime.now(MSK).isoformat()


def parse_num(raw, default=None):
    text = str(raw or "").strip().replace(",", ".")
    if not text:
        return default
    try:
        return float(text)
    except ValueError:
        return default


DIMS_RE = re.compile(r"^\s*(\d+[.,]?\d*)\s*[x×хX*]\s*(\d+[.,]?\d*)\s*[x×хX*]\s*(\d+[.,]?\d*)\s*$")


def parse_liters(raw):
    """Литры числом или габаритами в сантиметрах: 8x8x120 → 7.68 л.

    Возвращает (литры, подпись габаритов). Подпись пустая, если ввели число.
    """
    text = str(raw or "").strip()
    if not text:
        return None, ""
    hit = DIMS_RE.match(text)
    if hit:
        a, b, c = (float(x.replace(",", ".")) for x in hit.groups())
        return round(a * b * c / 1000.0, 4), "%g×%g×%g" % (a, b, c)
    return parse_num(text), ""


def parse_file(name, data):
    name = (name or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        import io
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        bits = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                bits.append(" ".join("" if v is None else str(v) for v in row))
        return parse_barcodes("\n".join(bits))
    text = data.decode("utf-8-sig", errors="replace")
    return parse_barcodes(text)


def parse_barcodes(text):
    seen = []
    have = set()
    for hit in BARCODE_RE.findall(str(text or "")):
        code = hit.strip()
        key = code.upper() if code.upper().startswith("OZN") else "".join(ch for ch in code if ch.isdigit()) or code
        if key in have:
            continue
        have.add(key)
        seen.append(code)
    return seen


def lookup(client_id, barcode, article="", hits=None):
    if hits is None:
        hits = find_cache_group(client_id, barcode=barcode or None, article=article or None)
    if not hits:
        return None
    first = prefer_hit(hits)[0]
    subject = (first["subject"] if "subject" in first.keys() else "") or ""
    name = first["name"] or ""
    gtin = ""
    for hit in prefer_hit(hits):
        val = (hit["gtin"] if "gtin" in hit.keys() else "") or ""
        if val:
            gtin = val
            break
    if not gtin:
        from catalog_pull import real_gtin

        gtin = real_gtin(barcode) or real_gtin(first["ext_barcode"])
    kind = (first["tracking_type"] if "tracking_type" in first.keys() else "") or ""
    if not kind:
        need = first["need_kiz"] if "need_kiz" in first.keys() else None
        kind = kind_from_subject((subject + " " + name).strip(), need_kiz=None if need is None else bool(need))
    return {
        "hits": hits,
        "article": first["ext_article"] or "",
        "name": name,
        "marketplace": "+".join(sorted({h["marketplace"] for h in hits if h["marketplace"]})),
        "gtin": gtin or "",
        "tracking_type": kind,
        "subject": subject,
    }


def article_groups(client_id, article):
    """Карточки с этим артикулом, разложенные по штрихкодам.

    У WB один артикул носят все размеры товара, поэтому несколько штрихкодов на
    один артикул — это несколько разных товаров, а не одна карточка.
    """
    buckets = {}
    for hit in find_cache(client_id, article=str(article or "").strip()):
        buckets.setdefault(barcode_norm(hit["ext_barcode"]), []).append(hit)
    return buckets


def match_registry(client_id, barcode, article):
    """Строка реестра → карточки кабинетов. Возвращает (hits, спорно, пояснение).

    Штрихкод — надёжный ключ. Артикул берём только запасным и только когда он
    ведёт ровно на один штрихкод: иначе гадать нельзя, отдаём строку оператору.
    """
    if barcode:
        hits = find_cache_group(client_id, barcode=barcode)
        if hits:
            return hits, "", ""
    article = str(article or "").strip()
    if not article:
        return [], "", ""
    buckets = article_groups(client_id, article)
    codes = [code for code in buckets if code]
    if len(codes) > 1:
        return (
            [hit for code in codes for hit in buckets[code]],
            "артикул %s ведёт на разные товары — выбери нужный" % article,
            "",
        )
    if codes:
        return find_cache_group(client_id, barcode=codes[0]), "", "сопоставлено по артикулу %s" % article
    flat = [hit for group in buckets.values() for hit in group]
    if flat:
        return flat, "", "сопоставлено по артикулу %s, штрихкода в кабинете нет" % article
    return [], "", ""


def group_norms(client_id, barcode, article=""):
    hits = find_cache_group(client_id, barcode=barcode or None, article=article or None)
    return {barcode_norm(h["ext_barcode"]) for h in hits if h["ext_barcode"]}, (hits[0]["gtin"] if hits and hits[0]["gtin"] else "")


def find_queued_same(client_id, barcode, gtin, article=""):
    """Уже в очереди тот же товар: тот же штрихкод, GTIN, артикул или карточка WB+Ozon."""
    norm = barcode_norm(barcode)
    norms, _g = group_norms(client_id, barcode)
    norms.add(norm)
    art = str(article or "").strip().upper()
    for row in list_intake(states=QUEUE_STATES):
        if row["client_id"] != client_id:
            continue
        if norm and barcode_norm(row["barcode"]) in norms:
            return row
        if gtin and row["gtin"] and row["gtin"] == gtin:
            return row
        # без штрихкода единственная зацепка — артикул клиента
        if not norm and art and str(row["article"] or "").strip().upper() == art:
            return row
        qnorm, _ = group_norms(client_id, row["barcode"], row["article"] or "")
        if norm and norm in qnorm:
            return row
    return None


def need_of(liters, tariff, qty):
    need = []
    if liters is None:
        need.append("литраж")
    if tariff is None:
        need.append("тариф")
    if not qty or float(qty) <= 0:
        need.append("количество")
    return need


def decide(found, liters, kind, gtin, tariff=None, qty=None):
    if not found:
        return "warn", "нет в кабинетах этого клиента"
    need = need_of(liters, tariff, qty)
    if need:
        return "warn", "нужно: " + ", ".join(need)
    if not kind:
        return "warn", "укажи тип продукции"
    if not gtin and tracking_code(kind) != "NOT_TRACKED":
        return "warn", "нужен GTIN"
    return "draft", ""


def fields_of(client_id, barcode, liters=None, kind="", gtin="", tariff=None, qty=None, dims="", hits=None, fallback=None):
    """fallback — артикул и наименование от клиента: нужны, когда кабинеты товар не знают."""
    fb = fallback or {}
    found = lookup(client_id, barcode, article=fb.get("article") or "", hits=hits)
    if not found:
        found = {"article": "", "name": "", "marketplace": "", "gtin": gtin14(barcode), "tracking_type": "", "hits": []}
    kind = (kind or "").strip() or found.get("tracking_type") or ""
    gtin = str(gtin or "").strip() or found.get("gtin") or ""
    state, note = decide(found if found.get("hits") else None, liters, kind, gtin, tariff, qty)
    return {
        "barcode": barcode,
        "article": found.get("article") or fb.get("article") or "",
        "name": found.get("name") or fb.get("name") or "",
        "marketplace": found.get("marketplace") or "",
        "gtin": gtin,
        "tracking_type": kind,
        "liters": liters,
        "dims": dims or "",
        "tariff": tariff,
        "qty": qty or 0,
        "state": state,
        "note": note,
    }


def refresh_client_catalog(client_id):
    from catalog_pull import pull_one

    notes = []
    for cab in list_cabinets():
        if cab["client_id"] != client_id or not cab["active"]:
            continue
        try:
            n, err = pull_one(cab)
            notes.append("%s %s" % (cab["marketplace"], err or ("%s товаров" % n)))
        except Exception as exc:
            notes.append("%s: %s" % (cab["marketplace"], exc))
    return notes


def add_many(client_id, barcodes, liters=None, kind="", author="", refresh=True, tariff=None, qty=None):
    client = get_client_by_id(int(client_id))
    if not client:
        raise ValueError("клиент не найден")
    codes = parse_barcodes("\n".join(barcodes) if not isinstance(barcodes, str) else barcodes)
    if isinstance(barcodes, (list, tuple)) and not codes:
        codes = [str(x).strip() for x in barcodes if str(x).strip()]
    if not codes:
        raise ValueError("нет штрихкодов")
    pulled = []
    if refresh:
        pulled = refresh_client_catalog(client["id"])
    added = []
    skipped = []
    found = 0
    missing = 0
    liters_val, dims_val = parse_liters(liters)
    qty_val = parse_num(qty)
    tariff_val = parse_num(tariff)
    if tariff_val is None:
        tariff_val = parse_num(client["tariff_storage"])
    for code in codes:
        fields = fields_of(
            client["id"], code, liters_val, kind, tariff=tariff_val, qty=qty_val, dims=dims_val
        )
        same = find_queued_same(client["id"], code, fields.get("gtin") or "")
        if same:
            mps = "+".join(sorted({p for p in ((same["marketplace"] or "") + "+" + (fields.get("marketplace") or "")).split("+") if p}))
            if mps and mps != (same["marketplace"] or ""):
                update_intake_row(same["id"], marketplace=mps)
            skipped.append({"barcode": code, "note": "тот же товар уже в очереди"})
            continue
        rid = add_intake_row(client["id"], fields, now_iso(), author)
        row = {"id": rid, "barcode": code, "state": fields["state"], "note": fields["note"], "name": fields["name"]}
        added.append(row)
        if fields["note"] == "нет в кабинетах этого клиента":
            missing += 1
        else:
            found += 1
    return {
        "added": added,
        "skipped": skipped,
        "found": found,
        "missing": missing,
        "pulled": pulled,
    }


def add_registry(client_id, filename, data, author="", refresh=True, tariff=None):
    """Реестр клиента xlsx → строки очереди, привязанные к карточке поставки."""
    client = get_client_by_id(int(client_id))
    if not client:
        raise ValueError("клиент не найден")
    parsed = registry.parse(filename, data)
    pulled = refresh_client_catalog(client["id"]) if refresh else []
    tariff_val = parse_num(tariff)
    if tariff_val is None:
        tariff_val = parse_num(client["tariff_storage"])
    supply_id = insert_supply(
        client["id"],
        parsed["supply"],
        filename or "",
        len(parsed["rows"]),
        parsed["qty_total"],
        now_iso(),
        author,
    )
    added = []
    skipped = []
    clashes = 0
    matched = 0
    for item in parsed["rows"]:
        hits, clash, hint = match_registry(client["id"], item["barcode"], item["article"])
        same = find_queued_same(client["id"], item["barcode"], "", item["article"])
        if same:
            skipped.append(
                {
                    "article": item["article"],
                    "barcode": item["barcode"],
                    "note": "тот же товар уже в очереди, строка %s" % same["id"],
                }
            )
            continue
        fields = fields_of(
            client["id"],
            item["barcode"],
            tariff=tariff_val,
            qty=item["qty"],
            hits=hits or None,
            fallback={"article": item["article"], "name": item["name"]},
        )
        if clash:
            fields["state"] = "clash"
            fields["note"] = clash
            clashes += 1
        elif hits:
            matched += 1
            if hint and fields["note"]:
                fields["note"] = "%s · %s" % (hint, fields["note"])
            elif hint:
                fields["note"] = hint
        rid = add_intake_row(client["id"], fields, now_iso(), author, supply_id=supply_id)
        added.append(
            {
                "id": rid,
                "article": item["article"],
                "barcode": item["barcode"],
                "qty": item["qty"],
                "state": fields["state"],
                "note": fields["note"],
            }
        )
    if not added:
        # весь реестр оказался дублем: карточка поставки без строк только путает
        delete_supply(supply_id)
        supply_id = None
    return {
        "supply_id": supply_id,
        "supply": parsed["supply"],
        "added": added,
        "skipped": skipped,
        "matched": matched,
        "clashes": clashes,
        "missing": sum(1 for r in added if r["note"] == "нет в кабинетах этого клиента"),
        "merged": parsed["merged"],
        "problems": parsed["problems"],
        "pulled": pulled,
    }


def candidates(row_id):
    """Варианты товара для спорной строки: по артикулу клиента, по одному на штрихкод."""
    row = get_intake_row(row_id)
    if not row:
        raise ValueError("строка не найдена")
    out = []
    for code, group in article_groups(row["client_id"], row["article"]).items():
        first = prefer_hit(group)[0]
        out.append(
            {
                "barcode": first["ext_barcode"] or code,
                "name": first["name"] or "",
                "size": (first["size"] if "size" in first.keys() else "") or "",
                "marketplace": "+".join(sorted({h["marketplace"] for h in group if h["marketplace"]})),
                "gtin": (first["gtin"] if "gtin" in first.keys() else "") or "",
            }
        )
    return sorted(out, key=lambda x: (x["size"], x["barcode"]))


def resolve(row_id, barcode):
    """Оператор выбрал товар для спорной строки: фиксируем штрихкод и пересчитываем."""
    row = get_intake_row(row_id)
    if not row:
        raise ValueError("строка не найдена")
    code = str(barcode or "").strip()
    if not code:
        raise ValueError("не выбран товар")
    if not find_cache_group(row["client_id"], barcode=code):
        raise ValueError("штрихкод %s не найден в кабинетах клиента" % code)
    fields = fields_of(
        row["client_id"],
        code,
        row["liters"],
        row["tracking_type"] or "",
        "",
        row["tariff"] if "tariff" in row.keys() else None,
        row["qty"],
        (row["dims"] if "dims" in row.keys() else "") or "",
        fallback={"article": row["article"], "name": row["name"]},
    )
    update_intake_row(
        row_id,
        barcode=code,
        article=fields["article"],
        name=fields["name"],
        marketplace=fields["marketplace"],
        gtin=fields["gtin"],
        tracking_type=fields["tracking_type"],
        state=fields["state"],
        note=fields["note"],
    )
    return get_intake_row(row_id)


def add(client_id, barcode, qty, liters, kind, gtin="", author=""):
    res = add_many(client_id, [barcode], liters=liters, kind=kind, author=author, refresh=False, qty=qty)
    if res["skipped"]:
        raise ValueError(res["skipped"][0]["note"])
    if not res["added"]:
        raise ValueError("не добавил строку")
    if gtin:
        patch(res["added"][0]["id"], gtin=gtin)
    return res["added"][0]["id"]


def patch(row_id, liters=None, kind=None, gtin=None, tariff=None, qty=None):
    row = get_intake_row(row_id)
    if not row:
        raise ValueError("строка не найдена")
    old_dims = (row["dims"] if "dims" in row.keys() else "") or ""
    if liters is None:
        liters_val, dims_val = row["liters"], old_dims
    else:
        liters_val, dims_val = parse_liters(liters)
    kind_val = row["tracking_type"] if kind is None else (kind or "").strip()
    gtin_val = row["gtin"] if gtin is None else str(gtin or "").strip()
    tariff_val = row["tariff"] if tariff is None else parse_num(tariff)
    qty_val = row["qty"] if qty is None else parse_num(qty)
    fields = fields_of(
        row["client_id"], row["barcode"], liters_val, kind_val, gtin_val, tariff_val, qty_val, dims_val,
        fallback={"article": row["article"], "name": row["name"]},
    )
    if row["state"] == "clash":
        # литраж и тариф править можно, но товар всё ещё не выбран
        fields["state"] = "clash"
        fields["note"] = row["note"]
    update_intake_row(
        row_id,
        article=fields["article"],
        name=fields["name"],
        marketplace=fields["marketplace"],
        gtin=fields["gtin"],
        tracking_type=fields["tracking_type"],
        liters=fields["liters"],
        dims=fields["dims"],
        tariff=fields["tariff"],
        qty=fields["qty"],
        state=fields["state"],
        note=fields["note"],
    )
    return get_intake_row(row_id)


def recheck(row):
    liters = row["liters"]
    tariff = row["tariff"] if "tariff" in row.keys() else None
    kind = row["tracking_type"] or ""
    need = need_of(liters, tariff, row["qty"])
    if need:
        return None, "нужно: " + ", ".join(need)
    found = lookup(row["client_id"], row["barcode"], row["article"] or "")
    if not found:
        return None, "нет в кабинетах этого клиента"
    gtin = (row["gtin"] or "").strip() or found["gtin"]
    kind = kind or found["tracking_type"]
    if not kind:
        return None, "укажи тип продукции"
    if not gtin and tracking_code(kind) != "NOT_TRACKED":
        return None, "нужен GTIN"
    first = prefer_hit(found["hits"])[0]
    payload = {
        "cabinet_id": first["cabinet_id"],
        "ext_key": first["ext_key"],
        "ext_article": first["ext_article"],
        "ext_barcode": row["barcode"] or first["ext_barcode"],
        "name": first["name"],
        "size": first["size"],
        "Литраж_л": liters,
        "Тип продукции": kind,
        "GTIN": gtin,
    }
    return {"payload": payload, "hits": found["hits"], "gtin": gtin, "marketplace": found["marketplace"]}, ""


def process_row(row):
    prepared, err = recheck(row)
    if err:
        update_intake_row(row["id"], state="warn", note=err)
        return None, err
    res = push_one(prepared["payload"], hits=prepared["hits"], dry=False)
    if not res["ok"]:
        update_intake_row(row["id"], state="warn", note=res["msg"])
        return None, res["msg"]
    for hit in prepared["hits"]:
        upsert_sku(
            row["client_id"],
            hit["cabinet_id"],
            hit["marketplace"],
            hit["ext_key"],
            hit["ext_article"],
            hit["ext_barcode"],
            res["id"],
        )
    update_intake_row(
        row["id"],
        article=res["article"],
        name=prepared["payload"]["name"],
        marketplace=prepared["marketplace"],
        gtin=prepared["gtin"],
        ms_product_id=res["id"],
        note=res["msg"],
    )
    return {"client": res["client"], "product_id": res["id"], "article": res["article"]}, ""


def run(blocking=True):
    init_db()
    with run_lock(blocking=blocking):
        return _run()


def _run():
    rows = list_intake(states=("draft", "warn"))
    done = []
    skipped = []
    pending = {}
    for row in rows:
        extra, err = process_row(row)
        if err:
            skipped.append({"id": row["id"], "barcode": row["barcode"], "note": err})
            print("приёмка %s: %s" % (row["barcode"], err))
            continue
        bucket = pending.setdefault(
            extra["client"]["id"], {"client": extra["client"], "items": [], "rows": []}
        )
        bucket["items"].append({"product_id": extra["product_id"], "qty": float(row["qty"])})
        bucket["rows"].append((row, extra))
    orders = []
    for bucket in pending.values():
        oid, number, err = create_purchase_order(bucket["client"], bucket["items"])
        if err:
            for row, _extra in bucket["rows"]:
                update_intake_row(row["id"], state="warn", note=err)
                skipped.append({"id": row["id"], "barcode": row["barcode"], "note": err})
            print("заказ поставщика: %s" % err)
            continue
        label = number or oid
        orders.append({"client": bucket["client"]["name"], "number": label, "ms_id": oid})
        for row, extra in bucket["rows"]:
            insert_lot(
                row["client_id"],
                extra["product_id"],
                extra["article"],
                row["barcode"],
                row["gtin"],
                row["name"],
                row["tracking_type"],
                row["liters"],
                float(row["qty"]),
                now_iso(),
                oid,
                row["tariff"] if "tariff" in row.keys() else None,
                row["dims"] if "dims" in row.keys() else "",
            )
            update_intake_row(row["id"], state="done", ms_order_name=label, note="заказ поставщика %s" % label)
            done.append({"id": row["id"], "barcode": row["barcode"], "order": label})
            print("приёмка %s: заказ поставщика %s" % (row["barcode"], label))
    return {"done": done, "skipped": skipped, "orders": orders}


if __name__ == "__main__":
    print(run())
